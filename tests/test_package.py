import hashlib
import importlib.util
import json
import os
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
PYTHON = os.environ.get("TEST_PYTHON", sys.executable)
FIXTURES = ROOT / "tests/fixtures"


def run(*args, check=True, env=None):
    return subprocess.run([str(x) for x in args], cwd=ROOT, text=True, capture_output=True, check=check, env=env)


class PackageTests(unittest.TestCase):
    def test_route_matrix(self):
        cases = {
            "做全年学业规划陪跑方案": ("jizhi-academic-year-plan-proposal", "single"),
            "做纯DP安心包方案": ("dp-proposal-designer", "single"),
            "做DP产品价格": ("dp-product-new-customer-quote", "single"),
            "做安心包报价": ("dp-product-new-customer-quote", "single"),
            "做DP全包作业方案并报价": ("dp-proposal-designer", "dp_proposal_with_quote"),
            "两门DP，一门陪跑并报价": ("jizhi-academic-year-plan-proposal", "mixed")
        }
        for text, expected in cases.items():
            result = json.loads(run(PYTHON, "scripts/route_request.py", text).stdout)
            self.assertEqual((result["owner"], result["mode"]), expected, text)

    def test_annual_deterministic(self):
        with tempfile.TemporaryDirectory() as tmp:
            a, b = Path(tmp)/"a.html", Path(tmp)/"b.html"
            script = "skills/jizhi-academic-year-plan-proposal/scripts/build_planning_proposal.py"
            run(PYTHON, script, FIXTURES/"annual.json", a)
            run(PYTHON, script, FIXTURES/"annual.json", b)
            self.assertEqual(hashlib.sha256(a.read_bytes()).digest(), hashlib.sha256(b.read_bytes()).digest())
            text = a.read_text(encoding="utf-8")
            for label in ("六大学业规划模块", "课程与服务匹配", "AI智慧学习系统", "每日", "每周", "每月", "方案价值", "学业规划价值", "AI智学系统价值", "押题价值", "陪跑课价值", "专业课价值"):
                self.assertIn(label, text)
            for internal in ("来源、边界", "预警提示", "报价追溯", "审核状态", "亲爱的学业规划师"):
                self.assertNotIn(internal, text)
            self.assertTrue(Path(str(a) + ".internal.html").exists())
            self.assertTrue(Path(str(a) + ".internal.json").exists())

    def test_annual_contracts_are_distinct(self):
        with tempfile.TemporaryDirectory() as tmp:
            base = json.loads((FIXTURES/"annual.json").read_text(encoding="utf-8"))
            expected = {
                "T00": "个性化学业规划报告",
                "T01": "AI智慧学习系统",
                "T02": "课程考核与服务安排",
                "T03": "服务报价",
                "T05": "DP与学业规划服务分工",
            }
            outputs = []
            for contract, marker in expected.items():
                data = dict(base)
                data["contract"] = contract
                if contract in {"T03", "T05"}:
                    data["quote"] = {"trace_id": "trace-test", "review_status": "已审核", "lines": [], "original_total": 100, "final_total": 100}
                source = Path(tmp)/f"{contract}.json"
                output = Path(tmp)/f"{contract}.html"
                source.write_text(json.dumps(data, ensure_ascii=False), encoding="utf-8")
                run(PYTHON, "skills/jizhi-academic-year-plan-proposal/scripts/build_planning_proposal.py", source, output)
                text = output.read_text(encoding="utf-8")
                self.assertIn(marker, text)
                self.assertIn("课程与服务匹配", text)
                self.assertNotIn("审核状态", text)
                self.assertNotIn("预警提示", text)
                self.assertNotIn('class="contract-label"', text)
                contracts = json.loads((ROOT/"skills/jizhi-academic-year-plan-proposal/assets/templates/template_contracts.json").read_text(encoding="utf-8"))
                self.assertIn(f'data-fixed-template="{contracts[contract]["fixed_case"]}"', text)
                if contract in {"T00", "T01", "T05"}:
                    self.assertIn('class="ip-hero"', text)
                else:
                    self.assertNotIn('class="ip-hero"', text)
                internal = Path(str(output) + ".internal.html").read_text(encoding="utf-8")
                self.assertIn("内部审核附件", internal)
                outputs.append(hashlib.sha256(output.read_bytes()).hexdigest())
            self.assertEqual(len(set(outputs)), len(expected))

    def test_t01_preflight_requires_visual_review(self):
        with tempfile.TemporaryDirectory() as tmp:
            html = Path(tmp) / "t01.html"
            pdf = Path(tmp) / "t01.pdf"
            fixed_pdf = ROOT / "skills/jizhi-academic-year-plan-proposal/templates/fixed_cases/固定模板01_标准年度学业规划方案.pdf"
            pdf.write_bytes(fixed_pdf.read_bytes())
            run(PYTHON, "skills/jizhi-academic-year-plan-proposal/scripts/build_planning_proposal.py", FIXTURES/"annual.json", html)
            script = "skills/jizhi-academic-year-plan-proposal/scripts/preflight_pdf.py"
            pending = run(PYTHON, script, html, pdf, check=False)
            self.assertEqual(pending.returncode, 2)
            self.assertFalse(json.loads(pdf.with_suffix(".preflight.json").read_text(encoding="utf-8"))["preflight_pass"])
            passed = run(PYTHON, script, html, pdf, "--visual-reviewed")
            result = json.loads(passed.stdout)
            self.assertTrue(result["preflight_pass"])
            self.assertEqual(result["contract"], "T01")
            self.assertEqual(result["fixed_case"], "固定模板01")
            self.assertEqual(result["pages"], 3)
            self.assertIn("已通过交付门禁", result["acceptance_declaration"])

    def test_dp_proposal_deterministic(self):
        with tempfile.TemporaryDirectory() as tmp:
            a, b = Path(tmp)/"a.html", Path(tmp)/"b.html"
            script = "skills/dp-proposal-designer/scripts/build_dp_proposal.py"
            run(PYTHON, script, FIXTURES/"dp_proposal.json", a)
            run(PYTHON, script, FIXTURES/"dp_proposal.json", b)
            self.assertEqual(a.read_bytes(), b.read_bytes())
            text = a.read_text(encoding="utf-8")
            self.assertNotIn("v2 官网检索版", text)
            self.assertIn("课程与服务匹配", text)
            self.assertIn("DP安心包核心价值", text)
            self.assertNotIn("预警提示", text)
            self.assertNotIn("报价状态", text)
            self.assertIn('class="ip-hero"', text)
            self.assertIn('class="contract-D01"', text)
            self.assertIn('data-fixed-template="固定模板06"', text)
            self.assertTrue(Path(str(a) + ".internal.html").exists())
            self.assertIn("DP内部审核附件", Path(str(a) + ".internal.html").read_text(encoding="utf-8"))

    def test_dp_mixed_module_contract(self):
        with tempfile.TemporaryDirectory() as tmp:
            data = json.loads((FIXTURES/"dp_proposal.json").read_text(encoding="utf-8"))
            data["contract"] = "D02"
            source, output = Path(tmp)/"d02.json", Path(tmp)/"d02.html"
            source.write_text(json.dumps(data, ensure_ascii=False), encoding="utf-8")
            run(PYTHON, "skills/dp-proposal-designer/scripts/build_dp_proposal.py", source, output)
            text = output.read_text(encoding="utf-8")
            self.assertIn("课程与服务匹配", text)
            self.assertNotIn("你的情况", text)
            self.assertNotIn('class="ip-hero"', text)

    def test_quote_workbook(self):
        with tempfile.TemporaryDirectory() as tmp:
            output = Path(tmp)/"quote.xlsx"
            run(PYTHON, "skills/dp-product-new-customer-quote/scripts/build_quote_workbook.py", FIXTURES/"dp_assessment.json", output, "--quote-json", FIXTURES/"dp_quote.json", "--reviewer", "QA")
            from openpyxl import load_workbook
            wb = load_workbook(output, data_only=False)
            self.assertEqual(wb.sheetnames, ["01课程考核汇总", "02课程与服务匹配", "03报价明细", "04资料来源", "05内部审核"])
            ws = wb["01课程考核汇总"]
            values = [cell.value for sheet in wb.worksheets for row in sheet.iter_rows() for cell in row]
            self.assertIn(12000, values)
            self.assertIn("折后价", values)
            self.assertEqual(ws["M10"].hyperlink.target, "https://example.edu/module")
            self.assertTrue(output.with_suffix(".xlsx.trace.json").exists())
            preflight = json.loads(output.with_suffix(".xlsx.preflight.json").read_text(encoding="utf-8"))
            self.assertTrue(preflight["preflight_pass"])
            self.assertEqual(preflight["fixed_case"], "固定模板04")

    def test_install_quarantines_retired_skill(self):
        with tempfile.TemporaryDirectory() as tmp:
            home = Path(tmp)
            for name in ("dp-customer-visual-proposal", "jizhi-academic-planning-report", "jizhi-essay-customer-proposal"):
                retired = home/f"skills/{name}"
                retired.mkdir(parents=True)
                (retired/"SKILL.md").write_text("old", encoding="utf-8")
            env = os.environ.copy(); env["CODEX_HOME"] = str(home); env["JIZHI_SKIP_RUNTIME_SETUP"] = "1"
            result = run("bash", "install.sh", env=env)
            self.assertIn("Quarantined retired competing skill", result.stdout)
            self.assertEqual(len(list((home/"skills").glob("*/SKILL.md"))), 3)

            output = home / "smoke"
            output.mkdir()
            run(PYTHON, home/"skills/jizhi-academic-year-plan-proposal/scripts/build_planning_proposal.py", FIXTURES/"annual.json", output/"annual.html")
            run(PYTHON, home/"skills/dp-proposal-designer/scripts/build_dp_proposal.py", FIXTURES/"dp_proposal.json", output/"dp.html")
            run(
                PYTHON,
                home/"skills/dp-product-new-customer-quote/scripts/build_quote_workbook.py",
                FIXTURES/"dp_assessment.json",
                output/"quote.xlsx",
                "--quote-json",
                FIXTURES/"dp_quote.json",
                "--reviewer",
                "QA",
            )
            for artifact in (output/"annual.html", output/"dp.html", output/"quote.xlsx"):
                self.assertGreater(artifact.stat().st_size, 0, artifact.name)

    def test_audit(self):
        result = run(PYTHON, "scripts/audit_boundaries.py")
        self.assertIn("PASS", result.stdout)


if __name__ == "__main__":
    unittest.main(verbosity=2)
