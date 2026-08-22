#!/usr/bin/env python3
"""Validate Assessment rows, optionally request a quote, and build one-sheet XLSX."""

from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path

FIELDS = ("school", "program", "degree_level", "target_year", "scope", "assessments")
ASSESSMENT_FIELDS = ("course_code", "course_name", "term", "assessment", "weight", "official_workload", "equivalent_words", "source_url", "evidence_status")


def validate(data: dict) -> list[str]:
    errors = [f"missing:{key}" for key in FIELDS if data.get(key) in (None, "", [])]
    for index, row in enumerate(data.get("assessments", [])):
        for key in ASSESSMENT_FIELDS:
            if row.get(key) in (None, ""):
                errors.append(f"assessments[{index}].missing:{key}")
        try:
            if int(row.get("equivalent_words", 0)) <= 0:
                errors.append(f"assessments[{index}].invalid:equivalent_words")
        except (TypeError, ValueError):
            errors.append(f"assessments[{index}].invalid:equivalent_words")
    return errors


def public_quote(quote: dict) -> dict:
    if quote.get("final_price") in (None, ""):
        raise ValueError("quote missing final_price")
    return {key: quote[key] for key in ("package_type", "scope", "original_price", "discount_rate", "final_price", "currency") if key in quote}


def build(data: dict, quote: dict, output: Path) -> None:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
    except ImportError as exc:
        raise SystemExit("openpyxl is required: python3 -m pip install openpyxl") from exc

    wb = Workbook()
    ws = wb.active
    ws.title = "课程考核与DP报价"
    navy, blue, mint, line = "0B2E63", "005CFF", "DFF8F3", "C9D8EA"
    ws.merge_cells("A1:N1")
    ws["A1"] = "课程 Assessment 与 DP 报价单"
    ws["A1"].font = Font(size=18, bold=True, color="FFFFFF")
    ws["A1"].fill = PatternFill("solid", fgColor=navy)
    ws["A1"].alignment = Alignment(horizontal="center")

    info = (("学校", data["school"]), ("专业", data["program"]), ("学历", data["degree_level"]), ("目标学年", data["target_year"]), ("覆盖范围", data["scope"]))
    for r, (label, value) in enumerate(info, 3):
        ws.cell(r, 1, label).font = Font(bold=True, color=navy)
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=6)
        ws.cell(r, 2, value)

    headers = ("序号", "课程代码", "课程名称", "学期", "Assessment", "占比", "官方工作量", "报价工作量", "要求", "风险", "判断依据", "证据状态", "信息源", "备注")
    header_row = 9
    for c, value in enumerate(headers, 1):
        cell = ws.cell(header_row, c, value)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=blue)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for i, row in enumerate(data["assessments"], header_row + 1):
        values = (i-header_row, row["course_code"], row["course_name"], row["term"], row["assessment"], row["weight"], row["official_workload"], int(row["equivalent_words"]), row.get("requirement", ""), row.get("risk", ""), row.get("evidence", ""), row["evidence_status"], "来源", row.get("note", ""))
        for c, value in enumerate(values, 1):
            cell = ws.cell(i, c, value)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        ws.cell(i, 13).hyperlink = row["source_url"]
        ws.cell(i, 13).style = "Hyperlink"

    quote_row = header_row + len(data["assessments"]) + 2
    total_words = sum(int(x["equivalent_words"]) for x in data["assessments"])
    quote_lines = (
        ("报价工作量合计", total_words),
        ("套餐", quote.get("package_type", data.get("package_type", "待确认"))),
        ("原价", quote.get("original_price", "待授权报价")),
        ("折扣", quote.get("discount_rate", "待授权报价")),
        ("折后价", quote.get("final_price", "待授权报价")),
        ("币种", quote.get("currency", "CNY")),
    )
    for offset, (label, value) in enumerate(quote_lines):
        r = quote_row + offset
        ws.cell(r, 11, label).font = Font(bold=True, color=navy)
        ws.cell(r, 12, value)
        ws.merge_cells(start_row=r, start_column=12, end_row=r, end_column=14)
        for c in range(11, 15):
            ws.cell(r, c).fill = PatternFill("solid", fgColor=mint)

    widths = [7, 14, 24, 12, 24, 10, 18, 14, 28, 20, 24, 14, 12, 22]
    for i, width in enumerate(widths, 1):
        ws.column_dimensions[chr(64+i)].width = width
    thin = Side(style="thin", color=line)
    for row in ws.iter_rows(min_row=header_row, max_row=header_row+len(data["assessments"]), min_col=1, max_col=14):
        for cell in row:
            cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)
    ws.freeze_panes = "A10"
    ws.auto_filter.ref = f"A9:N{header_row+len(data['assessments'])}"
    ws.sheet_view.showGridLines = False
    wb.save(output)


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("input_json")
    parser.add_argument("output_xlsx")
    parser.add_argument("--quote-json")
    parser.add_argument("--request-quote", action="store_true")
    args = parser.parse_args()
    data = json.loads(Path(args.input_json).read_text(encoding="utf-8"))
    errors = validate(data)
    if errors:
        raise SystemExit("Invalid assessment input: " + ", ".join(errors))
    quote = {}
    if args.quote_json:
        quote = public_quote(json.loads(Path(args.quote_json).read_text(encoding="utf-8")))
    elif args.request_quote:
        sys.path.insert(0, str(Path(__file__).resolve().parent))
        from quote_dp_api import request_quote
        payload = {key: data[key] for key in ("school", "program", "degree_level", "target_year", "scope")}
        payload["total_words"] = sum(int(x["equivalent_words"]) for x in data["assessments"])
        for key in ("package_type", "target_score"):
            if data.get(key):
                payload[key] = data[key]
        quote = request_quote(payload)
    build(data, quote, Path(args.output_xlsx))
    print(args.output_xlsx)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
