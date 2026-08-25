#!/usr/bin/env python3
"""Validate the fixed T04 DP assessment and quote workbook contract."""

import argparse
import json
from pathlib import Path

SHEETS = ["01课程考核汇总", "02课程与服务匹配", "03报价明细", "04资料来源", "05内部审核"]
HEADERS = {
    "01课程考核汇总": ["序号", "课程代码", "课程名称", "学期", "Assessment", "占比", "官方工作量", "报价工作量", "工作量依据", "审核状态", "证据状态", "来源年份", "信息源", "备注"],
    "02课程与服务匹配": ["课程名称", "课程代码", "课程考核形式", "课程工作量", "匹配服务"],
    "04资料来源": ["课程代码", "课程名称", "来源年份", "证据状态", "核验日期", "官网链接"],
}


def row_values(ws, row, width):
    return [ws.cell(row, column).value for column in range(1, width + 1)]


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("input_xlsx")
    args = parser.parse_args()
    from openpyxl import load_workbook

    path = Path(args.input_xlsx).resolve()
    workbook = load_workbook(path, data_only=False)
    checks = {
        "workbook_nonempty": path.exists() and path.stat().st_size >= 1000,
        "sheet_order": workbook.sheetnames == SHEETS,
        "summary_headers": row_values(workbook[SHEETS[0]], 9, len(HEADERS[SHEETS[0]])) == HEADERS[SHEETS[0]],
        "service_headers": row_values(workbook[SHEETS[1]], 3, len(HEADERS[SHEETS[1]])) == HEADERS[SHEETS[1]],
        "source_headers": row_values(workbook[SHEETS[3]], 3, len(HEADERS[SHEETS[3]])) == HEADERS[SHEETS[3]],
        "fixed_titles": [workbook[name]["A1"].value for name in SHEETS] == ["课程 Assessment 汇总", "课程与服务匹配", "DP报价明细", "资料来源", "内部审核与预警"],
        "fixed_style": all(workbook[name]["A1"].fill.fgColor.rgb in {"000B2E63", "0B2E63"} for name in SHEETS),
    }
    result = {
        "preflight_pass": all(checks.values()),
        "contract": "T04",
        "fixed_case": "固定模板04",
        "output": "xlsx",
        "checks": checks,
        "acceptance_declaration": "已通过交付门禁：contract=T04; fixed_case=固定模板04; sheets=5; layout=fixed; wording=fixed。" if all(checks.values()) else "未完成：交付门禁未全部通过。",
    }
    output = path.with_suffix(path.suffix + ".preflight.json")
    output.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(result, ensure_ascii=False, sort_keys=True))
    raise SystemExit(0 if result["preflight_pass"] else 1)


if __name__ == "__main__":
    main()
