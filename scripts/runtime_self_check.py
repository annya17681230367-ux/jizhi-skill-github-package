#!/usr/bin/env python3
"""Verify the isolated Jizhi runtime can create XLSX and PDF artifacts."""

import argparse
import subprocess
import tempfile
from pathlib import Path

from openpyxl import Workbook, load_workbook
from pypdf import PdfReader


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--browser", required=True)
    args = parser.parse_args()
    browser = Path(args.browser)
    if not browser.exists():
        raise SystemExit(f"Configured Chromium executable does not exist: {browser}")

    with tempfile.TemporaryDirectory() as tmp:
        root = Path(tmp)
        xlsx = root / "runtime-check.xlsx"
        workbook = Workbook()
        workbook.active["A1"] = "Excel运行时正常"
        workbook.save(xlsx)
        if load_workbook(xlsx, read_only=True).active["A1"].value != "Excel运行时正常":
            raise SystemExit("XLSX runtime verification failed.")

        html = root / "runtime-check.html"
        pdf = root / "runtime-check.pdf"
        html.write_text("<meta charset='utf-8'><h1>PDF运行时正常</h1>", encoding="utf-8")
        command = [
            str(browser), "--headless", "--disable-gpu", "--no-pdf-header-footer",
            f"--print-to-pdf={pdf}", html.resolve().as_uri(),
        ]
        subprocess.run(command, check=True, stdout=subprocess.PIPE, stderr=subprocess.PIPE)
        if not pdf.exists() or len(PdfReader(str(pdf)).pages) != 1:
            raise SystemExit("PDF runtime verification failed.")
    print("JIZHI_RUNTIME_SELF_CHECK=PASS")


if __name__ == "__main__":
    main()
